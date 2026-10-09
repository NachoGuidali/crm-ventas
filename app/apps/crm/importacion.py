"""
Importación masiva de prospectos (Excel / CSV).

Se procesa en segundo plano (Celery) fila por fila: cada fila pasa por
`ingresar_prospecto`, así que la deduplicación, la asignación rotativa y los
mensajes de bienvenida funcionan igual que en cualquier otro canal. Una fila con
error no frena al resto: queda registrada en el detalle del lote.
"""
import csv
import io
import logging
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.utils import timezone

logger = logging.getLogger('apps.crm')

CAMPOS_IMPORTACION = [
    ('', '— Ignorar columna —'),
    ('nombre', 'Nombre completo'),
    ('nombre_pila', 'Nombre (solo nombre)'),
    ('apellido', 'Apellido'),
    ('telefono', 'Teléfono'),
    ('telefono_alternativo', 'Teléfono alternativo'),
    ('email', 'Email'),
    ('dni', 'DNI'),
    ('fecha_atencion', 'Fecha de atención'),
    ('especialidad_atencion', 'Motivo / especialidad'),
    ('fecha_nacimiento', 'Fecha de nacimiento'),
    ('localidad', 'Localidad'),
    ('provincia', 'Provincia'),
    ('valor', 'Valor / cuota'),
    ('pauta', 'Pauta / origen de campaña'),
    ('nota', 'Nota (queda en la ficha)'),
    ('extra', 'Dato extra (se guarda tal cual)'),
]

_SUGERENCIAS = {
    'nombre': ['nombre y apellido', 'nombre completo', 'apellido y nombre', 'paciente', 'cliente', 'razon social'],
    'nombre_pila': ['nombre', 'nombres', 'first name'],
    'apellido': ['apellido', 'apellidos', 'last name'],
    'telefono': ['telefono', 'celular', 'movil', 'tel', 'whatsapp', 'phone', 'cel', 'telefono celular'],
    'telefono_alternativo': ['telefono 2', 'telefono alternativo', 'tel 2', 'otro telefono', 'telefono fijo'],
    'email': ['email', 'mail', 'correo', 'e-mail', 'correo electronico'],
    'dni': ['dni', 'documento', 'nro documento', 'doc'],
    'fecha_atencion': ['fecha de atencion', 'fecha atencion', 'fecha turno', 'fecha consulta', 'atencion'],
    'especialidad_atencion': ['especialidad', 'motivo', 'motivo consulta', 'servicio'],
    'fecha_nacimiento': ['fecha de nacimiento', 'nacimiento', 'fecha nac'],
    'localidad': ['localidad', 'ciudad', 'barrio'],
    'provincia': ['provincia', 'estado'],
    'valor': ['valor', 'cuota', 'monto', 'importe'],
    'nota': ['nota', 'notas', 'observaciones', 'comentario', 'comentarios'],
    'pauta': ['pauta', 'campana', 'campaign', 'utm campaign', 'origen', 'anuncio', 'ad name'],
}


def campos_importacion():
    """Campos fijos + campos personalizados activos (clave "cp:<slug>")."""
    from .models import CampoPersonalizado
    return CAMPOS_IMPORTACION + [(f'cp:{c.slug}', f'★ {c.nombre}') for c in CampoPersonalizado.activos() if not c.es_archivo]


def _norm(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9 ]+', ' ', texto).strip()


def sugerir_mapeo(columnas):
    from .models import CampoPersonalizado
    personalizados = {_norm(c.nombre): f'cp:{c.slug}' for c in CampoPersonalizado.activos() if not c.es_archivo}
    personalizados.update({_norm(c.slug.replace('_', ' ')): f'cp:{c.slug}'
                           for c in CampoPersonalizado.activos() if not c.es_archivo})
    mapeo, usados = {}, set()
    for col in columnas:
        n = _norm(col)
        if n in personalizados and personalizados[n] not in usados:
            mapeo[col] = personalizados[n]
            usados.add(personalizados[n])
            continue
        elegido = ''
        for campo, alias in _SUGERENCIAS.items():
            if campo in usados:
                continue
            if n in alias:
                elegido = campo
                break
        if not elegido:
            for campo, alias in _SUGERENCIAS.items():
                if campo not in usados and any(a in n for a in alias if len(a) > 3):
                    elegido = campo
                    break
        if elegido:
            usados.add(elegido)
        mapeo[col] = elegido or 'extra'
    return mapeo


def leer_archivo(archivo, nombre, limite=None):
    """Devuelve (columnas, filas: list[dict]). Soporta .xlsx y .csv (separador , o ; autodetectado)."""
    nombre = nombre.lower()
    if nombre.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        wb = load_workbook(archivo, read_only=True, data_only=True)
        ws = wb.active
        filas_iter = ws.iter_rows(values_only=True)
        encabezado = next(filas_iter, None) or []
        columnas = [str(c).strip() if c is not None else f'Columna {i + 1}' for i, c in enumerate(encabezado)]
        filas = []
        for fila in filas_iter:
            if fila is None or all(v in (None, '') for v in fila):
                continue
            filas.append({columnas[i]: v for i, v in enumerate(fila) if i < len(columnas)})
            if limite and len(filas) >= limite:
                break
        wb.close()
        return columnas, filas
    crudo = archivo.read()
    if isinstance(crudo, bytes):
        for enc in ('utf-8-sig', 'latin-1'):
            try:
                crudo = crudo.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    muestra = crudo[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=',;\t|')
    except csv.Error:
        dialecto = csv.excel
    reader = csv.DictReader(io.StringIO(crudo), dialect=dialecto)
    columnas = [c.strip() for c in (reader.fieldnames or [])]
    filas = []
    for fila in reader:
        if not any((v or '').strip() for v in fila.values() if isinstance(v, str)):
            continue
        filas.append({(k or '').strip(): v for k, v in fila.items()})
        if limite and len(filas) >= limite:
            break
    return columnas, filas


def _fecha(valor):
    if valor in (None, ''):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and 20000 < valor < 80000:  # número de serie de Excel
        from datetime import timedelta
        return date(1899, 12, 30) + timedelta(days=int(valor))
    texto = str(valor).strip().split(' ')[0]
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d.%m.%Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'fecha no reconocida: "{valor}"')


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _decimal(valor):
    """Acepta 1500 · 1500.5 · 1.500,50 · $ 1,500.50"""
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = re.sub(r'[^\d.,-]', '', str(valor))
    if ',' in texto and '.' in texto:
        texto = texto.replace('.', '').replace(',', '.') if texto.rfind(',') > texto.rfind('.') else texto.replace(',', '')
    elif ',' in texto:
        texto = texto.replace(',', '.') if len(texto.split(',')[-1]) <= 2 else texto.replace(',', '')
    elif texto.count('.') > 1 or (texto.count('.') == 1 and len(texto.split('.')[-1]) == 3):
        texto = texto.replace('.', '')
    try:
        return Decimal(texto) if texto else None
    except InvalidOperation:
        return None


def _valor_personalizado(campo, valor):
    T = campo.__class__
    if campo.tipo == T.TIPO_FECHA:
        f = _fecha(valor)
        return f.isoformat() if f else None
    if campo.tipo == T.TIPO_NUMERO:
        d = _decimal(valor)
        return (int(d) if d == int(d) else float(d)) if d is not None else None
    if campo.tipo == T.TIPO_SINO:
        return _norm(valor) in ('si', 's', 'true', '1', 'x', 'yes')
    if campo.tipo == T.TIPO_TELEFONO:
        from core.phone import normalizar_telefono
        return normalizar_telefono(_texto(valor)) or _texto(valor)
    if campo.tipo == T.TIPO_LISTA:
        texto = _texto(valor)
        coincide = next((o for o in campo.opciones if _norm(o) == _norm(texto)), None)
        if coincide is None:
            raise ValueError(f'"{texto}" no es una opción válida de {campo.nombre}')
        return coincide
    return _texto(valor)


def fila_a_datos(fila, mapeo, personalizados=None):
    datos, extra, notas = {}, {}, []
    nombre_pila = apellido = ''
    for col, campo in mapeo.items():
        valor = fila.get(col)
        if not campo or valor in (None, ''):
            continue
        if campo in ('fecha_atencion', 'fecha_nacimiento'):
            datos[campo] = _fecha(valor)
        elif campo == 'valor':
            datos['valor'] = _decimal(valor)
        elif campo == 'nombre_pila':
            nombre_pila = _texto(valor)
        elif campo == 'apellido':
            apellido = _texto(valor)
        elif campo == 'nota':
            notas.append(_texto(valor))
        elif campo == 'pauta':
            datos['origen_pauta'] = _texto(valor)[:200]
        elif campo == 'extra':
            extra[col] = _texto(valor)
        elif campo.startswith('cp:'):
            cp = (personalizados or {}).get(campo[3:])
            if cp is not None:
                extra[cp.slug] = _valor_personalizado(cp, valor)
        elif campo == 'dni':
            datos['dni'] = re.sub(r'\D', '', _texto(valor))[:12]
        else:
            datos[campo] = _texto(valor)
    if not datos.get('nombre') and (nombre_pila or apellido):
        datos['nombre'] = f'{nombre_pila} {apellido}'.strip()
    if datos.get('nombre'):
        datos['nombre'] = datos['nombre'].title() if datos['nombre'].isupper() else datos['nombre']
    if extra:
        datos['datos_extra'] = extra
    return datos, notas


def procesar_lote(lote_id):
    from .models import Actividad, ImportacionLote, Oportunidad
    from .services import ErrorNegocio, ingresar_prospecto

    lote = ImportacionLote.objects.select_related('embudo', 'etapa', 'agente_fijo', 'creado_por').get(pk=lote_id)
    if lote.estado not in (ImportacionLote.ESTADO_PENDIENTE, ImportacionLote.ESTADO_ERROR):
        return 'ya procesado'
    ImportacionLote.objects.filter(pk=lote.pk).update(estado=ImportacionLote.ESTADO_PROCESANDO)
    try:
        with lote.archivo.open('rb') as f:
            _, filas = leer_archivo(f, lote.nombre_archivo)
    except Exception as e:
        ImportacionLote.objects.filter(pk=lote.pk).update(estado=ImportacionLote.ESTADO_ERROR,
                                                          errores_detalle=[{'fila': 0, 'error': f'No se pudo leer: {e}'}])
        return 'error de lectura'

    from .models import CampoPersonalizado
    personalizados = {c.slug: c for c in CampoPersonalizado.activos() if not c.es_archivo}
    contadores = {'procesados': 0, 'creados': 0, 'contactos_nuevos': 0, 'ya_existentes': 0, 'errores': 0}
    errores = []
    ImportacionLote.objects.filter(pk=lote.pk).update(total=len(filas))

    for n, fila in enumerate(filas, start=2):  # fila 1 = encabezado
        try:
            datos, notas = fila_a_datos(fila, lote.mapeo, personalizados)
            if not (datos.get('telefono') or datos.get('email') or datos.get('dni')):
                raise ErrorNegocio('sin teléfono, email ni DNI')
            valor = datos.pop('valor', None)
            origen_pauta = datos.pop('origen_pauta', '')
            res = ingresar_prospecto(
                datos, lote.embudo, Oportunidad.ORIGEN_IMPORTACION, fuente=lote.fuente or lote.nombre_archivo,
                usuario=lote.creado_por, etapa=lote.etapa, lote=lote, agente=lote.agente_fijo,
                asignar=lote.asignar_automaticamente, disparar_automatizaciones=lote.disparar_automatizaciones,
                valor=valor, origen_pauta=origen_pauta, pauta=None if origen_pauta else lote.pauta,
            )
            if res.oportunidad_nueva:
                contadores['creados'] += 1
            else:
                contadores['ya_existentes'] += 1
            if res.contacto_nuevo:
                contadores['contactos_nuevos'] += 1
            for nota in notas:
                Actividad.objects.create(contacto=res.contacto, oportunidad=res.oportunidad, tipo=Actividad.TIPO_NOTA,
                                         usuario=lote.creado_por, texto=f'[Importación] {nota}')
        except Exception as e:
            contadores['errores'] += 1
            if len(errores) < 500:
                errores.append({'fila': n, 'error': str(e)[:200],
                                'datos': {k: _texto(v)[:40] for k, v in list(fila.items())[:4]}})
            if not isinstance(e, (ErrorNegocio, ValueError)):
                logger.exception('Error importando fila %s del lote %s', n, lote.pk)
        contadores['procesados'] += 1
        if contadores['procesados'] % 25 == 0:
            ImportacionLote.objects.filter(pk=lote.pk).update(**contadores)

    ImportacionLote.objects.filter(pk=lote.pk).update(
        **contadores, errores_detalle=errores, estado=ImportacionLote.ESTADO_COMPLETADO, finalizado_at=timezone.now(),
    )
    if lote.creado_por:
        from apps.users.services import notificar
        notificar(lote.creado_por, 'sistema', f'Importación terminada: {lote.nombre_archivo}',
                  f'{contadores["creados"]} nuevos · {contadores["ya_existentes"]} ya existían · '
                  f'{contadores["errores"]} con error', f'/importaciones/{lote.pk}/')
    return contadores
